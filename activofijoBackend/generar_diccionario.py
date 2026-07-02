import os
import django
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "settings")
django.setup()

from django.apps import apps

def set_cell_background(cell, fill, color, val):
    # This sets the cell background color
    shading_elm = parse_xml(r'<w:shd {} w:fill="{}" w:color="{}" w:val="{}"/>'.format(nsdecls('w'), fill, color, val))
    cell._tc.get_or_add_tcPr().append(shading_elm)

doc = Document()

title = doc.add_heading('DICCIONARIO DE BASE DE DATOS', level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

app_models = apps.get_app_config('activo').get_models()
all_models = list(app_models)

# Definir prioridad de modelos
priority = {
    'in_activo': 1,
    'in_asignado': 2,
    'in_det_asig': 3,
    'in_ingreso': 4,
    'in_baja_act': 5,
    'in_revaluo': 6,
    'in_oficina': 7,
    'in_unidad': 8,
    'in_grupo': 9,
    'in_responsable': 10,
    'in_gestion': 11,
}

all_models.sort(key=lambda x: priority.get(x.__name__, 999))

for model in all_models:
    # Title for the table
    heading = doc.add_heading(f'TABLA {model.__name__.upper()}', level=2)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    # Create table
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    headers = ['CAMPO', 'DOMINIO', 'TIPO DE CAMPO', 'INDICACION', 'VALIDACION']
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        # Set header background to a strong blue like the image
        set_cell_background(hdr_cells[i], "5B9BD5", "auto", "clear")
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.bold = True
    
    for row_idx, field in enumerate(model._meta.fields):
        row_cells = table.add_row().cells
        
        # Alternating background colors like the image
        bg_color = "D2DEEF" if row_idx % 2 == 0 else "EAEFF7"
        for cell in row_cells:
            set_cell_background(cell, bg_color, "auto", "clear")
            
        # 1. CAMPO
        row_cells[0].text = field.name
        row_cells[0].paragraphs[0].runs[0].font.bold = True
        
        # 2. DOMINIO
        row_cells[1].text = str(field.verbose_name).capitalize()
        
        # 3. TIPO DE CAMPO
        internal_type = field.get_internal_type()
        tipo = "String"
        if internal_type in ['IntegerField', 'AutoField', 'SmallIntegerField', 'BigIntegerField']:
            tipo = "Integer"
        elif internal_type in ['FloatField', 'DecimalField']:
            tipo = "Float"
        elif internal_type in ['DateField', 'DateTimeField']:
            tipo = "Date"
        elif internal_type == 'BooleanField':
            tipo = "Boolean"
        elif internal_type == 'ForeignKey':
            tipo = "Integer (FK)"
        row_cells[2].text = tipo
        
        # 4. INDICACION
        indicacion = "Texto"
        if tipo == "Integer":
            indicacion = "Número entero"
        elif tipo == "Float":
            indicacion = "Número con decimales"
        elif tipo == "Date":
            indicacion = "Fecha/Hora"
        elif tipo == "Boolean":
            indicacion = "Verdadero o Falso"
            
        if hasattr(field, 'max_length') and field.max_length:
            indicacion += f" (máx. {field.max_length} car.)"
            
        if internal_type == 'ForeignKey':
            indicacion = f"Clave de tabla {field.related_model.__name__}"
            
        row_cells[3].text = indicacion
        
        # 5. VALIDACION
        validaciones = []
        if field.primary_key:
            validaciones.append("Clave Primaria")
        if not field.null and not field.blank:
            validaciones.append("Entrada obligatoria")
        else:
            validaciones.append("Entrada opcional")
            
        if field.unique:
            validaciones.append("única")
            
        row_cells[4].text = " y ".join(validaciones)

    doc.add_paragraph() # Add some space

output_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'Diccionario_Base_Datos.docx')
doc.save(output_path)
print(f"Documento generado exitosamente en: {output_path}")
